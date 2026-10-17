"""Gera amostras ANONIMIZADAS dos relatórios de renovação de cada seguradora.

Mesma estrutura dos arquivos reais (formato real, aba, cabeçalho, ordem e formato das
colunas), com 5 linhas de dados TOTALMENTE INVENTADOS. Saída:
    tests/fixtures/planilhas/<seguradora>/exemplo-anonimizado.<extensão original>

Uso: python scripts/gerar_fixtures_seguradoras.py
Layouts documentados em docs/importacao/.
"""

from __future__ import annotations

import random
from pathlib import Path

import openpyxl
import xlwt

RAIZ = Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "tests" / "fixtures" / "planilhas"
rnd = random.Random(2026)

NOMES = ["ANA FICTICIA DOS SANTOS", "BRUNO EXEMPLO LIMA", "CARLA INVENTADA ROCHA", "DIEGO TESTE MOURA", "PADARIA MODELO LTDA"]
NOMES_TITULO = [n.title().replace(" Dos ", " dos ").replace(" Ltda", " LTDA") for n in NOMES]
VENCIMENTOS = [(5, 11, 2026), (12, 11, 2026), (18, 11, 2026), (24, 11, 2026), (30, 11, 2026)]


def cpf() -> str:
    base = [rnd.randint(0, 9) for _ in range(9)]
    for pesos in (range(10, 1, -1), range(11, 1, -1)):
        s = sum(d * p for d, p in zip(base, pesos)) % 11
        base.append(0 if s < 2 else 11 - s)
    return "".join(map(str, base))


def cnpj() -> str:
    base = [rnd.randint(0, 9) for _ in range(8)] + [0, 0, 0, 1]
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        s = sum(d * p for d, p in zip(base, pesos)) % 11
        base.append(0 if s < 2 else 11 - s)
    return "".join(map(str, base))


def doc(i: int) -> str:
    return cnpj() if "LTDA" in NOMES[i] else cpf()


def num(n: int) -> str:
    return "".join(str(rnd.randint(0, 9)) for _ in range(n))


def br(v: float, casas: int = 2) -> str:
    """1234.5 -> '1.234,50' (milhar com ponto)."""
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def br_sem_milhar(v: float, casas: int = 2) -> str:
    return f"{v:.{casas}f}".replace(".", ",")


def dmy(d: tuple[int, int, int], ano_delta: int = 0) -> str:
    return f"{d[0]:02d}/{d[1]:02d}/{d[2] + ano_delta}"


def ymd(d: tuple[int, int, int], ano_delta: int = 0) -> str:
    return f"{d[2] + ano_delta}-{d[1]:02d}-{d[0]:02d}"


def email(i: int) -> str:
    return f"cliente{i + 1}@exemplo.com"


def placa() -> str:
    return "TST" + str(rnd.randint(0, 9)) + "A" + num(2)


def chassi() -> str:
    return "9TESTE" + num(11)


def pasta(seg: str) -> Path:
    p = DESTINO / seg
    p.mkdir(parents=True, exist_ok=True)
    return p


# --------------------------------------------------------------------------- Allianz: .xls binário (BIFF)
def allianz() -> None:
    cab = ["Apólice", "Item", "CPF", "Nome", "Ramo", "Fim de Vigencia", "Prazo para Renovação", "Prêmio Liq. Atual", "Prêmio Liq. Anterior", "Status"]
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("title.aplication")
    for j, c in enumerate(cab):
        ws.write(0, j, c)
    for i, v in enumerate(VENCIMENTOS):
        prazo = (v[0], v[1] - 1, v[2])  # prazo para renovar: antes do vencimento
        linha = [num(9) + "J" + num(9), "1", cpf(), NOMES[i], "1211", dmy(v), dmy(prazo), "0", br(rnd.uniform(800, 3500)), "   "]
        for j, val in enumerate(linha):
            ws.write(i + 1, j, val)  # tudo como texto, como no arquivo real
    wb.save(str(pasta("allianz") / "exemplo-anonimizado.xls"))


# --------------------------------------------------------------------------- HDI: texto separado por TAB com extensão .xls
def hdi() -> None:
    cab = ["Empresa", "Sucursal", "Carteira", "TipDocum", "SeqDocum", "Endosso", "Nome", "CliEspecial", "FimVigencia", "PremioLiq", "TotIndeniza"]
    linhas = ["\t".join(cab)]
    for i, v in enumerate(VENCIMENTOS):
        linhas.append("\t".join(["01", "068", "431", "A", num(6), num(6), NOMES[i], "ESPECIAL", dmy(v), br_sem_milhar(rnd.uniform(800, 3500), 4), ""]))
    # ASCII, CRLF entre linhas, sem quebra no final (como o arquivo real)
    (pasta("hdi") / "exemplo-anonimizado.xls").write_bytes("\r\n".join(linhas).encode("ascii"))


# --------------------------------------------------------------------------- Mapfre: CSV ';' em cp1252
def mapfre() -> None:
    cab = ["Apólice", "Segurado", "Fim de Vigência", "Produto", "Modalidade", "Cód. Interno", "Status", "Motivo"]
    linhas = [";".join(cab)]
    for i, v in enumerate(VENCIMENTOS):
        linhas.append(";".join([num(12), NOMES[i] + " ", ymd(v) + " 00:00:00.0", "231-MAPFRE AUTO", "Auto Compreensivo", num(5), "A renovar", ""]))
    (pasta("mapfre") / "exemplo-anonimizado.csv").write_bytes(("\n".join(linhas) + "\n").encode("cp1252"))


# --------------------------------------------------------------------------- Porto: .xlsx, tudo como texto
def porto() -> None:
    cab = ["nomeRazaoSocial", "nomeSocialPessoa", "nomeProduto", "codigoProduto", "codigoPessoa", "ramo", "sucursal", "apolice", "endosso",
           "origemProposta", "numeroProposta", "susep", "dataEmissao", "dataFinalVigencia", "dataInicioVigencia", "status", "cpfCnpj", "email", "celular", "ano"]
    produtos = [("PROT COMB AUTO+RESID", "68"), ("Automóvel", "27"), ("Azul Automóvel", "188"), ("PROT COMB AUTO+RESID", "68"), ("Automóvel", "27")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(cab)
    for i, v in enumerate(VENCIMENTOS):
        nome = NOMES_TITULO[i]
        cel = f"(11) 9XX{num(2)}-XX{num(2)}"  # a Porto já envia o celular parcialmente mascarado
        ws.append([nome, nome, produtos[i][0], produtos[i][1], num(8), "531", "22", num(8), "0", rnd.choice(["12", "20"]), num(8),
                   num(5) + "J", ymd((v[0] - 1, v[1] - 1, v[2]), -1), ymd(v), ymd(v, -1), "ATIVO", doc(i), email(i) if i != 2 else None, cel,
                   "25" if i % 2 else None])
    wb.save(pasta("porto") / "exemplo-anonimizado.xlsx")


# --------------------------------------------------------------------------- Suhai: tabela HTML (GridView) com extensão .xls
def suhai() -> None:
    cab = ["ID da Renovacao", "Status da Renovacao", "N&#250;mero da Ap&#243;lice", "Status da Ap&#243;lice", "In&#237;cio de Vig&#234;ncia",
           "T&#233;rmino de Vig&#234;ncia", "Dias de Vig&#234;ncia", "Cancelamento de Vig&#234;ncia", "Marca", "Modelo", "Placa", "Chassi",
           "Pr&#234;mio L&#237;quido", "Pr&#234;mio Anualizado", "Pr&#234;mio Pago", "Visilibidade", "ID do C&#225;lculo", "Nome do Proponente",
           "CPF/CNPJ do Proponente", "Telefone do Proponente", "Celular do Proponente", "Email do Proponente", "ID do Corretor",
           "CPF/CNPJ do Corretor", "C&#243;digo SUSEP do Corretor", "Nome do Corretor", "Email do Corretor", "Comiss&#227;o do Corretor",
           "Tabela de Comiss&#227;o", "ID da Assessoria", "Nome da Assessoria", "Fator de Ajuste FIPE", "Classe de B&#244;nus"]
    th = "".join(f'<th scope="col">{c}</th>' for c in cab)
    trs = []
    for i, v in enumerate(VENCIMENTOS):
        premio = rnd.uniform(60, 4000)
        valores = [num(7), "A Renovar", num(13), "Emitida", dmy(v, -1) + " 00:00:00", dmy(v) + " 00:00:00", "365", "&nbsp;",
                   "MARCA EXEMPLO", "MODELO TESTE 1.0", placa(), chassi(), br_sem_milhar(premio), br_sem_milhar(premio), br_sem_milhar(premio * 1.3),
                   num(22), "&nbsp;", NOMES[i], doc(i).lstrip("0"), f"1140{num(6)}" if i % 2 else "&nbsp;", f"1190{num(7)}", email(i),
                   "10001", "00000000000191", "100000001", "CORRETORA EXEMPLO LTDA", "contato@corretora-exemplo.com", rnd.choice(["16%", "20%"]),
                   rnd.choice(["tab 24", "tab 5"]), "50001", "ASSESSORIA EXEMPLO LTDA.", "10000", str(rnd.randint(0, 3))]
        tds = "".join(f"<td>{x}</td>" for x in valores)
        trs.append(f'\t\t<tr style="background-color:White;">\n\t\t\t{tds}\n\t\t</tr>')
    html = ('<div>\n\t<table cellspacing="0" rules="all" border="1" style="border-collapse:collapse;">\n'
            f'\t\t<tr style="color:White;background-color:DarkGray;">\n\t\t\t{th}\n\t\t</tr>' + "".join("\n" + t for t in trs) +
            "\n\t</table>\n</div>")
    (pasta("suhai") / "exemplo-anonimizado.xls").write_text(html, encoding="utf-8")


# --------------------------------------------------------------------------- Tokio: .xlsx com extensão .xls
def tokio() -> None:
    cab = ["Produto", "Código do Cliente", "Nome Cliente", "CPF/CNPJ Segurado", "CPF/CNPJ Condutor", "E-mail Segurado", "Motivo/Situação",
           "Apólice Anterior", "Negocio Anterior", "Item Anterior", "Início Vigência Anterior", "Fim Vigência Anterior", "Negócio Renovação",
           "Item Renovação", "Início Vigência Renovação", "Fim Vigência Renovação", "Código do Corretor", "Nome Empresa Parceira",
           "Código Empresa Parceira", "Código do Parceiro de Negocio", "Código Crivo", "Quantidade Sinistro", "Quantidade Parcelas",
           "Descrição Item Segurado", "Placa", "Chassi", "Cep Risco", "Bônus", "Código Sucursal", "Bônus Interno", "CAP", "Telefone de Contato",
           "Renovação Facilitada"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    ws.append(cab)
    for i, v in enumerate(VENCIMENTOS):
        documento = doc(i)
        bonus = str(rnd.randint(0, 10))
        ws.append([rnd.choice(["Auto", "Moto"]), num(8), NOMES[i], documento, documento, email(i).upper(), "Disponível para Cálculo",
                   num(8), num(9), num(9), dmy(v, -1), dmy(v), num(9), num(9), dmy(v), dmy(v, 1), "10001", None, None, None,
                   str(rnd.randint(0, 999)), "0", rnd.choice(["1", "5", "12"]), "MARCA-MODELO 1.0 FLEX", placa(), chassi(), "01000" + num(3),
                   bonus, "8106", bonus, str(rnd.randint(-9, 99)), f"(11) 4000{num(4)} / (11) 9000{num(5)}", "NÃO"])
    # o arquivo real é .xlsx com extensão .xls
    destino = pasta("tokio") / "exemplo-anonimizado.xls"
    wb.save(destino)


# --------------------------------------------------------------------------- Zurich: exportação vazia (só estilo, sem tabela)
def zurich() -> None:
    (pasta("zurich") / "exemplo-anonimizado.xls").write_text("<style> .text { mso-number-format:\\@; } </style><div>\r\n\r\n</div>", encoding="utf-8")


if __name__ == "__main__":
    for gerar in (allianz, hdi, mapfre, porto, suhai, tokio, zurich):
        gerar()
    print("Amostras geradas em", DESTINO)
